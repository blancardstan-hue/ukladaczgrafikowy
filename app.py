import streamlit as st
import pandas as pd
import random
import itertools
from io import BytesIO
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Border, Side

st.set_page_config(page_title="Układacz Grafików", layout="wide")

st.title("📅 Szkolny Układacz Grafików")
st.write("Witaj w aplikacji! Pobierz szablon, wypełnij go, a następnie wgraj poniżej.")

# ==========================================
# 1. GENERATOR SZABLONU EXCEL (UREALNIONY)
# ==========================================
def generate_excel_template():
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        
        imiona = [
            "Anna", "Maria", "Katarzyna", "Małgorzata", "Agnieszka", "Ewa", "Magdalena", 
            "Julia", "Zofia", "Hanna", "Jan", "Piotr", "Krzysztof", "Andrzej", "Tomasz", 
            "Paweł", "Michał", "Marcin", "Jakub", "Adam", "Stanisław", "Tymoteusz", "Juliana"
        ]
        nazwiska = [
            "Nowak", "Kowalski", "Wiśniewski", "Wójcik", "Kowalczyk", "Kamiński", "Lewandowski", 
            "Zieliński", "Szymański", "Woźniak", "Dąbrowski", "Kozłowski", "Jankowski", "Mazur", 
            "Kwiatkowski", "Krawczyk", "Kaczmarek", "Piotrowski", "Grabowski", "Zając", "Blancard"
        ]
        
        kombinacje = [f"{i} {n}" for i in imiona for n in nazwiska]
        wybrani = random.sample(kombinacje, 72)
        
        # Dominacja pełnych etatów i szerokich popołudni
        godziny_lek_opcje = [
            "08:00-21:00", "08:00-21:00", "08:00-21:00", "08:00-21:00", 
            "12:00-20:00", "13:00-21:00", "14:00-21:00",
            "15:00-20:00", "16:00-21:00", "08:00-15:00"
        ]
        
        regiony_filii = [
            ["Komorów", "Michałowice"], ["Komorów", "Michałowice"], ["Komorów"], ["Michałowice"],
            ["Ursus 1", "Ursus 2"], ["Ursus 1"], ["Ursus 2"],
            ["Pruszków", "Nowa Wieś"], ["Pruszków"], ["Nowa Wieś"],
            ["Komorów", "Pruszków"], ["Wszystkie"]
        ]
        
        poziomy_opcje = [
            "3-5 lat, 0, 1, 2, 3", 
            "4, 5, 6, 7, 8, Masters", 
            "0, 1, 2, 3, 4, 5, 6, 7, 8, Masters"
        ]
        
        lektorzy_data = []
        for nazwa in wybrani:
            # Przesunięcie wagi na osoby z 1-3 grupami (metodycy, liderzy itp.)
            min_gr = random.choices([1, 2, 3, 4, 5, 6, 7], weights=[15, 20, 20, 15, 15, 10, 5])[0]
            variance = random.choices([0, 1, 2], weights=[20, 60, 20])[0]
            max_gr = min_gr + variance
            if max_gr > 8: max_gr = 8
            
            pref = random.choices(["TAK", "NIE"], weights=[5, 95])[0]
            przypisane_filie = random.choice(regiony_filii)
            if "Wszystkie" in przypisane_filie:
                filie_str = "Komorów, Michałowice, Pruszków, Ursus 1, Ursus 2, Nowa Wieś"
            else:
                filie_str = ", ".join(przypisane_filie)
            
            lektorzy_data.append({
                "Lektor": nazwa,
                "Preferencyjne traktowanie": pref,
                "Min liczba grup": min_gr,
                "Max liczba grup": max_gr,
                "Dostępność Pon": random.choice(godziny_lek_opcje),
                "Dostępność Wt": random.choice(godziny_lek_opcje),
                "Dostępność Śr": random.choice(godziny_lek_opcje),
                "Dostępność Czw": random.choice(godziny_lek_opcje),
                "Dostępność Pt": random.choice(godziny_lek_opcje),
                "Filie": filie_str,
                "Poziomy": random.choice(poziomy_opcje)
            })
            
        pd.DataFrame(lektorzy_data).to_excel(writer, sheet_name="Lektorzy", index=False)
        
        filie_sale = {
            "Nowa Wieś": 2,
            "Komorów": 5,
            "Michałowice": 5,
            "Pruszków": 5,
            "Ursus 1": 7,
            "Ursus 2": 7
        }
        
        sale_data = []
        for f, cnt in filie_sale.items():
            for i in range(1, cnt + 1):
                prz = "3-5 lat, 0, 1, 2, 3" if i == 1 else "4, 5, 6, 7, 8, Masters" if i == 2 else "Wszystkie"
                sale_data.append({"Nazwa Sali": str(i), "Filia": f, "Przeznaczenie": prz})
        pd.DataFrame(sale_data).to_excel(writer, sheet_name="Sale", index=False)
        
        szkoly_config = {
            "Nowa Wieś": {"SP Nowa Wieś": 2}, 
            "Komorów": {"SP Komorów": 4},     
            "Michałowice": {"SP Michałowice": 4},
            "Pruszków": {"SP nr 1 Pruszków": 4, "SP nr 2 Pruszków": 3},
            "Ursus 1": {"SP Ursus A": 5},     
            "Ursus 2": {"SP Ursus B": 5}
        }
        szkoly_lo = {"Komorów": ["LO Komorów"], "Pruszków": ["LO Pruszków"], "Nowa Wieś": ["LO Nowa Wieś"]}
        godziny_pocz = ["08:00", "08:55", "09:50"]
        godziny_kon = ["12:30", "13:30", "14:25", "15:20", "16:15"]

        mlodsze_data = []
        for f, szkoly in szkoly_config.items():
            for sz, oddzialy in szkoly.items():
                mlodsze_data.append({
                    "Nazwa Grupy": f"Zerówka {sz}", "Poziom": "0", 
                    "Liczba Dzieci": random.randint(8, 14),
                    "Czas trwania (min)": 60, "Skąd Odbiór": sz, "Docelowa Filia": f,
                    "Początek Szkoły Pon": "08:00", "Koniec Szkoły Pon": "12:30",
                    "Początek Szkoły Wt": "08:00", "Koniec Szkoły Wt": "12:30",
                    "Początek Szkoły Śr": "08:00", "Koniec Szkoły Śr": "12:30",
                    "Początek Szkoły Czw": "08:00", "Koniec Szkoły Czw": "13:30",
                    "Początek Szkoły Pt": "08:00", "Koniec Szkoły Pt": "12:30",
                    "Liczba Spotkań": 2
                })
        pd.DataFrame(mlodsze_data).to_excel(writer, sheet_name="Młodsze Dzieci", index=False)
        
        starsi_data = []
        litery_full = ['A', 'B', 'C', 'D', 'E']
        for f, szkoly in szkoly_config.items():
            for sz, oddzialy in szkoly.items():
                for kl in range(1, 9):
                    uzyte_oddzialy = oddzialy - 1 if kl == 6 else oddzialy
                    for lit in litery_full[:uzyte_oddzialy]:
                        starsi_data.append({
                            "Szkoła i Klasa": f"{sz} - Klasa {kl}{lit}", "Poziom": str(kl), 
                            "Liczba Chętnych": random.randint(4, 10), "Czas trwania (min)": 90, 
                            "Docelowa Filia": f,
                            "Początek Szkoły Pon": random.choice(godziny_pocz), "Koniec Szkoły Pon": random.choice(godziny_kon),
                            "Początek Szkoły Wt": random.choice(godziny_pocz), "Koniec Szkoły Wt": random.choice(godziny_kon),
                            "Początek Szkoły Śr": random.choice(godziny_pocz), "Koniec Szkoły Śr": random.choice(godziny_kon),
                            "Początek Szkoły Czw": random.choice(godziny_pocz), "Koniec Szkoły Czw": random.choice(godziny_kon),
                            "Początek Szkoły Pt": random.choice(godziny_pocz), "Koniec Szkoły Pt": random.choice(godziny_kon)
                        })
                        
        for f, szkoly in szkoly_lo.items():
            for sz in szkoly:
                for kl in range(1, 5):
                    for lit in ['A', 'B']:
                        starsi_data.append({
                            "Szkoła i Klasa": f"{sz} - Klasa {kl}{lit}", "Poziom": str(kl + 8), 
                            "Liczba Chętnych": random.randint(4, 10), "Czas trwania (min)": 90, 
                            "Docelowa Filia": f,
                            "Początek Szkoły Pon": "08:00", "Koniec Szkoły Pon": "15:20",
                            "Początek Szkoły Wt": "08:00", "Koniec Szkoły Wt": "16:15",
                            "Początek Szkoły Śr": "08:00", "Koniec Szkoły Śr": "14:25",
                            "Początek Szkoły Czw": "08:00", "Koniec Szkoły Czw": "15:20",
                            "Początek Szkoły Pt": "08:00", "Koniec Szkoły Pt": "13:30"
                        })
        pd.DataFrame(starsi_data).to_excel(writer, sheet_name="Starsi Uczniowie", index=False)
        
        pd.DataFrame({
            "Imię Opiekunki": ["Marta", "Krystyna", "Zofia", "Ewa", "Agnieszka", "Magda", "Joanna", "Monika"],
            "Filia": ["Komorów", "Michałowice", "Pruszków", "Ursus 1", "Ursus 2", "Nowa Wieś", "Ursus 1", "Komorów"],
            "Dostępność Pon": ["12:00-18:00"]*8, "Dostępność Wt": ["12:00-18:00"]*8,
            "Dostępność Śr": ["12:00-18:00"]*8, "Dostępność Czw": ["12:00-18:00"]*8,
            "Dostępność Pt": ["12:00-18:00"]*8
        }).to_excel(writer, sheet_name="Opiekunki", index=False)
        
        trasy_data = []
        for f, szkoly in szkoly_config.items():
            for sz in szkoly.keys(): 
                trasy_data.append({"Początek": sz, "Koniec": f, "Czas (min)": random.choice([10, 15, 20])})
        pd.DataFrame(trasy_data).to_excel(writer, sheet_name="Trasy", index=False)

        for sheetname in writer.sheets:
            ws = writer.sheets[sheetname]
            for i, col in enumerate(ws.columns, 1):
                max_len = max([len(str(c.value)) for c in col if c.value] + [0])
                ws.column_dimensions[get_column_letter(i)].width = max_len + 3
                
    return output.getvalue()

# ==========================================
# FUNKCJE POMOCNICZE
# ==========================================
dni_short = ["Pon", "Wt", "Śr", "Czw", "Pt"]
dni_pocz = ["Początek Szkoły Pon", "Początek Szkoły Wt", "Początek Szkoły Śr", "Początek Szkoły Czw", "Początek Szkoły Pt"]
dni_kon = ["Koniec Szkoły Pon", "Koniec Szkoły Wt", "Koniec Szkoły Śr", "Koniec Szkoły Czw", "Koniec Szkoły Pt"]

def time_to_mins(t_str):
    if pd.isna(t_str) or not str(t_str).strip(): return -1
    try:
        h, m = map(int, str(t_str).split(':'))
        return h * 60 + m
    except: return -1

def mins_to_time(mins):
    if mins < 0: return ""
    return f"{int(mins // 60):02d}:{int(mins % 60):02d}"

def parse_availability(avail_str):
    if pd.isna(avail_str) or not str(avail_str).strip(): return []
    blocks = []
    for part in str(avail_str).split(','):
        try:
            st_s, en_s = part.split('-')
            blocks.append((time_to_mins(st_s), time_to_mins(en_s)))
        except: pass
    return blocks

def is_overlap(st1, en1, st2, en2, gap=0):
    return max(st1, st2 - gap) < min(en1, en2 + gap)

def score_combo(combo):
    if len(combo) == 1: return 1
    if len(combo) == 2:
        diff = abs(dni_short.index(combo[0]) - dni_short.index(combo[1]))
        if diff in [2, 4]: return 1
        elif diff == 3: return 2
        elif diff == 1: return 3
        return 10
    return 1

def check_lektor(lek, d, st_m, en_m, grafik, gap, zad_filia, zakaz_migracji_aktywny):
    if zakaz_migracji_aktywny:
        dzisiejsze_filie = [g["Filia"] for g in grafik if g["Lektor"] == lek["Lektor"] and g["Dzień"] == d]
        if dzisiejsze_filie and zad_filia not in dzisiejsze_filie:
            return False

    can_work = any(b_s <= st_m and b_e >= en_m for (b_s, b_e) in lek["Avail"][d])
    if not can_work: return False
    zajety = any(
        g["Lektor"] == lek["Lektor"] and g["Dzień"] == d and 
        is_overlap(st_m, en_m, g["Start"], g["End"], gap) 
        for g in grafik
    )
    return not zajety

def check_sala(filia, poziom, d, st_m, en_m, grafik, sale_dane, gap, lektor):
    dostepne = [s["Sala"] for s in sale_dane if s["Filia"] == filia and (poziom in s["Poziomy"] or "Wszystkie" in s["Poziomy"])]
    sale_dzis = [g["Sala"] for g in grafik if g["Lektor"] == lektor and g["Dzień"] == d and g["Filia"] == filia]
    sale_ogolnie = [g["Sala"] for g in grafik if g["Lektor"] == lektor and g["Filia"] == filia]
    
    def sala_score(s):
        if s in sale_dzis: return 0
        if s in sale_ogolnie: return 1
        return 2
        
    dostepne.sort(key=sala_score)
    for s in dostepne:
        zajeta = any(g["Sala"] == s and g["Dzień"] == d and g["Filia"] == filia and is_overlap(st_m, en_m, g["Start"], g["End"], gap) for g in grafik)
        if not zajeta: return s
    return None

colors_pool = [
    "#FFB3BA", "#FFDFBA", "#FFFFBA", "#BAFFC9", "#BAE1FF",
    "#D5AAFF", "#FFB3E6", "#B3FFB3", "#FFC9DE", "#E0BBE4",
    "#957DAD", "#D291BC", "#FEC8D8", "#FFDFD3", "#F5D0C5"
]

def style_cell_with_cmap(val, cmap):
    if val == "-" or pd.isna(val) or val == "": return ""
    val_str = str(val)
    chunks = val_str.split(" | \n")
    visible_chunks = [c for c in chunks if not c.strip().startswith("~")]
    
    check_val = chunks[0].replace("~", "")
    bg_color = ""
    for k, color in cmap.items():
        if f"({k})" in check_val or check_val.startswith(k):
            bg_color = color
            break
            
    if not bg_color: return ""

    if not visible_chunks:
        return f"background-color: {bg_color}; color: {bg_color};"
    else:
        return f"background-color: {bg_color}; color: #000000; font-weight: bold;"

def style_cell_with_cmap_lektor(val, cmap):
    if val == "-" or pd.isna(val) or val == "": return ""
    val_str = str(val)
    chunks = val_str.split(" | \n")
    visible_chunks = [c for c in chunks if not c.strip().startswith("~")]
    
    check_val = chunks[0].replace("~", "")
    bg_color = ""
    for k, color in cmap.items():
        if k in check_val:
            bg_color = color
            break
            
    if not bg_color: return ""

    if not visible_chunks:
        return f"background-color: {bg_color}; color: {bg_color};"
    else:
        return f"background-color: {bg_color}; color: #000000; font-weight: bold;"

def format_cell_text(val):
    if val == "-" or pd.isna(val) or val == "": return val
    val_str = str(val)
    chunks = val_str.split(" | \n")
    visible_chunks = [c.replace("~", "") for c in chunks if not c.strip().startswith("~")]
    if not visible_chunks: return ""
    return " | \n".join(visible_chunks)

def apply_lesson_to_grid(df_grid, st_m, en_m, col_key, wpis):
    start_day = 480  
    end_day = 1260   
    time_intervals = [(m, m+30) for m in range(start_day, end_day, 30)]
    idx_labels = [mins_to_time(m) for m in range(start_day, end_day, 30)]

    intervals_touched = []
    for i, (win_st, win_en) in enumerate(time_intervals):
        overlap = min(en_m, win_en) - max(st_m, win_st)
        if overlap > 0:
            intervals_touched.append((i, overlap))
    
    if not intervals_touched: return
    
    num_blocks = max(1, round((en_m - st_m) / 30.0))
    best_window = []
    max_ov = -1
    
    for j in range(len(intervals_touched) - num_blocks + 1):
        window = intervals_touched[j : j + num_blocks]
        ov = sum(x[1] for x in window)
        if ov > max_ov:
            max_ov = ov
            best_window = window
            
    if not best_window:
        best_window = intervals_touched
        
    main_i = best_window[0][0]
    
    for i, overlap in best_window:
        current_wpis = wpis if i == main_i else "~" + wpis
        current_val = df_grid.at[idx_labels[i], col_key]
        if current_val: df_grid.at[idx_labels[i], col_key] = str(current_val) + " | \n" + current_wpis
        else: df_grid.at[idx_labels[i], col_key] = current_wpis

def generate_grid_for_filia(df_f, filia_nazwa, sale_dane):
    sale_w_filii = sorted([s["Sala"] for s in sale_dane if s["Filia"] == filia_nazwa])
    if not sale_w_filii: sale_w_filii = ["1", "2", "3", "4", "5"]
    
    multi_cols = pd.MultiIndex.from_product([dni_short, sale_w_filii], names=["Dzień", "Sala"])
    
    start_day = 480  
    end_day = 1260   
    idx_labels = [mins_to_time(m) for m in range(start_day, end_day, 30)]
    
    df_grid = pd.DataFrame(index=idx_labels, columns=multi_cols).fillna("")
    if df_f.empty: return df_grid
    
    for _, row in df_f.iterrows():
        d = row["Dzień"]
        s = row["Sala"]
        if d in dni_short and s in sale_w_filii:
            wpis = f"{mins_to_time(row['Start'])}-{mins_to_time(row['End'])}\n{row['Grupa']} ({row['Lektor']}){row.get('Op_Str', '')}"
            apply_lesson_to_grid(df_grid, row["Start"], row["End"], (d, s), wpis)
                        
    return df_grid

def generate_grid_for_lektor(df_lek):
    start_day = 480
    end_day = 1260
    idx_labels = [mins_to_time(m) for m in range(start_day, end_day, 30)]
    
    df_grid = pd.DataFrame(index=idx_labels, columns=dni_short).fillna("")
    if df_lek.empty: return df_grid
    
    for _, row in df_lek.iterrows():
        d = row["Dzień"]
        if d in dni_short:
            wpis = f"{mins_to_time(row['Start'])}-{mins_to_time(row['End'])}\n{row['Grupa']}\n📍 {row['Filia']} (Sala: {row['Sala']})"
            apply_lesson_to_grid(df_grid, row["Start"], row["End"], d, wpis)
                    
    return df_grid

def create_excel_download(grafiki_dict):
    output = BytesIO()
    thin = Side(style='thin', color='D3D3D3')
    medium = Side(style='medium', color='000000')

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for nazwa, styled_df in grafiki_dict.items():
            sheet_name = nazwa[:31] 
            styled_df.to_excel(writer, sheet_name=sheet_name)
            ws = writer.sheets[sheet_name]
            
            df = styled_df.data if hasattr(styled_df, "data") else styled_df
            
            for row in ws.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
                    if isinstance(cell.value, str):
                        chunks = cell.value.split(" | \n")
                        visible_chunks = [c.replace("~", "") for c in chunks if not c.strip().startswith("~")]
                        if visible_chunks: cell.value = " | \n".join(visible_chunks)
                        else: cell.value = ""

                    cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)

            for row in ws.iter_rows(min_col=1, max_col=1):
                for cell in row:
                    b = cell.border
                    new_top = medium if cell.row == 1 else b.top
                    new_bottom = medium if cell.row == ws.max_row else b.bottom
                    cell.border = Border(top=new_top, bottom=new_bottom, left=medium, right=medium)

            if isinstance(df.columns, pd.MultiIndex):
                for i in range(len(df.columns)):
                    current_day = df.columns[i][0]
                    is_first = (i == 0) or (df.columns[i-1][0] != current_day)
                    is_last = (i == len(df.columns) - 1) or (df.columns[i+1][0] != current_day)
                    excel_col = i + 2

                    for row in ws.iter_rows(min_col=excel_col, max_col=excel_col):
                        for cell in row:
                            b = cell.border
                            new_left = medium if is_first else b.left
                            new_right = medium if is_last else b.right
                            new_top = medium if cell.row == 1 else b.top
                            new_bottom = medium if cell.row == ws.max_row else b.bottom
                            cell.border = Border(top=new_top, bottom=new_bottom, left=new_left, right=new_right)
            else:
                for i in range(len(df.columns)):
                    excel_col = i + 2
                    for row in ws.iter_rows(min_col=excel_col, max_col=excel_col):
                        for cell in row:
                            b = cell.border
                            new_top = medium if cell.row == 1 else b.top
                            new_bottom = medium if cell.row == ws.max_row else b.bottom
                            cell.border = Border(top=new_top, bottom=new_bottom, left=medium, right=medium)

            for i in range(1, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(i)].width = 25
                
    return output.getvalue()

# ==========================================
# INTERFEJS GŁÓWNY
# ==========================================
st.subheader("1. Pobierz szablon")
st.download_button("📥 Pobierz układacz (Excel)", data=generate_excel_template(), file_name="ukladacz.xlsx")

st.subheader("2. Wgraj uzupełniony plik")
uploaded_file = st.file_uploader("Wgraj plik Excel", type=["xlsx"])

if uploaded_file is not None:
    xls = pd.ExcelFile(uploaded_file)
    df_lektorzy = pd.read_excel(xls, "Lektorzy")
    
    if "Preferencyjne traktowanie" not in df_lektorzy.columns:
        df_lektorzy.insert(1, "Preferencyjne traktowanie", "NIE")
    if "Min liczba grup" not in df_lektorzy.columns:
        df_lektorzy.insert(2, "Min liczba grup", 0)
    if "Max liczba grup" not in df_lektorzy.columns:
        if "Zadeklarowana liczba grup" in df_lektorzy.columns:
            df_lektorzy.insert(3, "Max liczba grup", df_lektorzy["Zadeklarowana liczba grup"])
        else:
            df_lektorzy.insert(3, "Max liczba grup", 100)
        
    df_sale = pd.read_excel(xls, "Sale")
    
    sale_dane = []
    for _, r in df_sale.iterrows():
        s_poz = [p.strip() for p in str(r["Przeznaczenie"]).split(",")] if pd.notna(r["Przeznaczenie"]) else []
        if "Masters" in s_poz: s_poz.extend(["9", "10", "11", "12"])
        sale_dane.append({"Sala": str(r["Nazwa Sali"]), "Filia": r["Filia"], "Poziomy": s_poz})
        
    df_trasy = pd.read_excel(xls, "Trasy")
    df_opiekunki = pd.read_excel(xls, "Opiekunki") if "Opiekunki" in xls.sheet_names else pd.DataFrame()
    
    df_mlodsze = pd.read_excel(xls, "Młodsze Dzieci")
    if "Czas trwania (min)" not in df_mlodsze.columns: df_mlodsze.insert(3, "Czas trwania (min)", 60)
    
    df_starsi = pd.read_excel(xls, "Starsi Uczniowie")
    if "Czas trwania (min)" not in df_starsi.columns: df_starsi.insert(4, "Czas trwania (min)", 90)

    st.header("Krok 1: Formowanie")
    filie_unikalne = sorted(df_sale["Filia"].dropna().unique())
    bufory_filii = {}
    cols_bufory = st.columns(3)
    for idx, f in enumerate(filie_unikalne):
        with cols_bufory[idx % 3]: 
            bufory_filii[f] = st.slider(f"{f} (min)", 15, 60, 30, 5, key=f"b_{f}")
            
    st.markdown("---")
    st.write("**Młodsze Dzieci**")
    edyt_mlodsze = st.data_editor(df_mlodsze, use_container_width=True, num_rows="dynamic", key="e_ml")

    st.write("**Starsi Uczniowie**")
    df_s = df_starsi[df_starsi["Liczba Chętnych"] > 0].copy()
    
    def ext_sch(x): return str(x).split(" - ")[0] if " - " in str(x) else x
    def ext_lit(x): 
        if " - " not in str(x): return ""
        part = str(x).split(" - ")[-1].replace("Klasa", "").replace(" ", "")
        return "".join([c for c in part if c.isalpha()])
        
    df_s["Szkoła"] = df_s["Szkoła i Klasa"].apply(ext_sch)
    df_s["Litera"] = df_s["Szkoła i Klasa"].apply(ext_lit)
    
    for d in dni_kon: df_s[d + "_mins"] = df_s[d].apply(time_to_mins)
    
    def avg_t(r):
        vals = [r[d+"_mins"] for d in dni_kon if r[d+"_mins"]>0]
        return sum(vals)/max(1, len(vals)) if vals else 0
        
    df_s["Avg_Time"] = df_s.apply(avg_t, axis=1)

    utworzone_g = []
    for _, g_df in df_s.groupby(["Docelowa Filia", "Szkoła", "Poziom"]):
        g_df = g_df.sort_values("Avg_Time")
        paczka, c_count = [], 0
        for _, row in g_df.iterrows():
            if c_count + row["Liczba Chętnych"] > 12:
                if c_count > 0: utworzone_g.append(paczka)
                paczka, c_count = [row], row["Liczba Chętnych"]
            else:
                paczka.append(row)
                c_count += row["Liczba Chętnych"]
        if paczka: utworzone_g.append(paczka)

    f_rows = []
    for p in utworzone_g:
        t_chet = sum(r["Liczba Chętnych"] for r in p)
        if t_chet < 5: continue
        
        litery = "".join(sorted([r['Litera'] for r in p]))
        r_data = {
            "Nazwa Grupy": f"{p[0]['Szkoła']} - {p[0]['Poziom']}{litery}",
            "Poziom": str(p[0]["Poziom"]), "Liczba Dzieci": t_chet,
            "Czas trwania (min)": max(r["Czas trwania (min)"] for r in p),
            "Docelowa Filia": p[0]["Docelowa Filia"], "Liczba Spotkań": 2 
        }
        for d_p, d_k in zip(dni_pocz, dni_kon):
            v_k = [r[d_k + "_mins"] for r in p if r[d_k + "_mins"] > 0]
            v_p = [time_to_mins(r[d_p]) for r in p if time_to_mins(r[d_p]) > 0]
            r_data[d_k] = mins_to_time(max(v_k)) if v_k else ""
            r_data[d_p] = mins_to_time(min(v_p)) if v_p else ""
        f_rows.append(r_data)

    df_aktywne = pd.DataFrame(f_rows)
    edyt_starsi = st.data_editor(df_aktywne, use_container_width=True, num_rows="dynamic", key="e_st")

    # ==========================================
    # KROK 2: GRAFIK (SESSION_STATE)
    # ==========================================
    st.markdown("---")
    st.header("Krok 2: Automatyczne Układanie Grafiku")
    
    zakaz_migracji_ui = st.checkbox("🚫 Zakaz migracji między filiami w obrębie jednego dnia", value=True)
    
    if "wygenerowano" not in st.session_state:
        st.session_state.wygenerowano = False
        
    if st.button("🚀 Wygeneruj Grafik", type="primary"):
        with st.spinner("Przeszukuję okna, zasoby i preferencje..."):
            grafik, grafik_op, nieprzypisane = [], [], []
            
            czas_tras = {f"{r['Początek']}_{r['Koniec']}": r["Czas (min)"] for _, r in df_trasy.iterrows()}
            
            lektorzy_dane = []
            for _, r in df_lektorzy.iterrows():
                l_poziomy = [p.strip() for p in str(r["Poziomy"]).split(",")] if pd.notna(r["Poziomy"]) else []
                if "Masters" in l_poziomy: l_poziomy.extend(["9", "10", "11", "12"])
                    
                l_filie = [f.strip() for f in str(r["Filie"]).split(",")] if pd.notna(r["Filie"]) else []
                avail = {dni_short[i]: parse_availability(r.get(f"Dostępność {dni_short[i]}", "")) for i in range(5)}
                
                pref_val = str(r.get("Preferencyjne traktowanie", "NIE")).strip().upper()
                is_pref = (pref_val == "TAK")
                
                lim_min = int(r.get("Min liczba grup", 0))
                lim_max = int(r.get("Max liczba grup", 100))
                
                lektorzy_dane.append({
                    "Lektor": r["Lektor"], "Poziomy": l_poziomy, "Filie": l_filie, 
                    "Avail": avail, "Limit_Min": lim_min, "Limit_Max": lim_max,
                    "Pref": is_pref
                })
                
            opiekunki_dane = []
            if not df_opiekunki.empty:
                for _, r in df_opiekunki.iterrows():
                    avail = {dni_short[i]: parse_availability(r.get(f"Dostępność {dni_short[i]}", "")) for i in range(5)}
                    opiekunki_dane.append({"Imię": r["Imię Opiekunki"], "Filia": r.get("Filia", ""), "Avail": avail})

            zadania = []
            def zbuduj_okna(r, typ):
                av_windows = {}
                do = str(r["Docelowa Filia"])
                bufor = bufory_filii.get(do, 30) if typ == "Starsza" else czas_tras.get(f"{r.get('Skąd Odbiór', '')}_{do}", 15) + 5
                
                for idx, (d_p, d_k) in enumerate(zip(dni_pocz, dni_kon)):
                    st_sz = time_to_mins(r.get(d_p, ""))
                    en_sz = time_to_mins(r.get(d_k, ""))
                    windows = []
                    if st_sz > 0: windows.append((480, max(480, st_sz - bufor)))
                    if en_sz > 0: windows.append((en_sz + bufor, 1050 if typ == "Młodsza" else 1230))
                    if windows: av_windows[dni_short[idx]] = windows
                return av_windows

            for _, r in edyt_mlodsze.iterrows():
                zadania.append({"Grupa": r["Nazwa Grupy"], "Poziom": str(r["Poziom"]), "Filia": str(r["Docelowa Filia"]), "Czas": r["Czas trwania (min)"], "Spotkań": int(r.get("Liczba Spotkań", 2)), "Typ": "Młodsza", "Odbiór": str(r.get("Skąd Odbiór", "")), "Windows": zbuduj_okna(r, "Młodsza")})
                
            for _, r in edyt_starsi.iterrows():
                zadania.append({"Grupa": r["Nazwa Grupy"], "Poziom": str(r["Poziom"]), "Filia": str(r["Docelowa Filia"]), "Czas": r["Czas trwania (min)"], "Spotkań": int(r.get("Liczba Spotkań", 2)), "Typ": "Starsza", "Odbiór": "", "Windows": zbuduj_okna(r, "Starsza")})
                
            zadania.sort(key=lambda x: (len(x["Windows"]), -x["Czas"]))
            
            tracker = {l["Lektor"]: {"grupy": set(), "poziomy": [], "filie_dni": {}} for l in lektorzy_dane}
            filia_day_load = {}
            
            def get_load(filia, days):
                f_load = filia_day_load.get(filia, {})
                return sum(f_load.get(d, 0) for d in days)
            
            for zad in zadania:
                valid_days = list(zad["Windows"].keys())
                if len(valid_days) >= zad["Spotkań"]:
                    combos = list(itertools.combinations(valid_days, zad["Spotkań"]))
                    random.shuffle(combos) 
                    combos.sort(key=lambda c: (score_combo(c), get_load(zad["Filia"], c)))
                else:
                    combos = []

                valid_leks = [l for l in lektorzy_dane if zad["Poziom"] in l["Poziomy"] and (not l["Filie"] or zad["Filia"] in l["Filie"])]
                waga = 0 if zad["Typ"] == "Młodsza" or zad["Poziom"] in ["0","1","2","3"] else 1 if zad["Poziom"] in ["4","5"] else 2
                
                # ZAMKNIĘCIE W FUNKCJĘ ABY WYKONAĆ 2 FAZY ALGORYTMU
                def proba_przypisania(strict_mode):
                    gap_preferences = [10, 5] if strict_mode else [5, 0]
                    
                    for gap_pref in gap_preferences:
                        for combo in combos:
                            def lek_score(lek):
                                trk = tracker[lek["Lektor"]]
                                score = 0
                                
                                hit_max = len(trk["grupy"]) >= lek["Limit_Max"]
                                hit_min = len(trk["grupy"]) >= lek["Limit_Min"]
                                
                                if hit_max: score += 10000
                                if hit_min: score += 100
                                
                                if strict_mode and lek["Pref"] and not hit_min:
                                    score -= 800
                                
                                for d in combo:
                                    dni_filie = trk["filie_dni"].get(d, set())
                                    if dni_filie:
                                        if zad["Filia"] not in dni_filie: 
                                            score += 2000 if strict_mode else 500  
                                        else: 
                                            score -= 60    
                                            
                                level_count = trk["poziomy"].count(zad["Poziom"])
                                
                                if strict_mode and lek["Pref"]:
                                    score -= (level_count * 50)
                                    if len(trk["poziomy"]) > 0 and level_count == 0:
                                        score += 30 # Drastycznie zmniejszona kara za nowy poziom
                                else:
                                    score -= (level_count * 15) 
                                
                                score += len(trk["grupy"]) 
                                return score

                            combo_valid_leks = sorted(valid_leks, key=lek_score)
                            
                            for lek in combo_valid_leks:
                                zaplanowane_dni = []
                                for d in combo:
                                    sloty = []
                                    for win_st, win_en in zad["Windows"][d]:
                                        st_test = win_st + (5 - win_st % 5) if win_st % 5 != 0 else win_st
                                        while st_test + zad["Czas"] <= win_en:
                                            sloty.append(st_test)
                                            st_test += 5
                                            
                                    if waga == 0: sloty.sort()
                                    elif waga == 2: sloty.sort(reverse=True)
                                    
                                    dzien_ok = False
                                    for st_m in sloty:
                                        en_m = st_m + zad["Czas"]
                                        
                                        aktywna_blokada = zakaz_migracji_ui and strict_mode
                                        if check_lektor(lek, d, st_m, en_m, grafik, gap_pref, zad["Filia"], aktywna_blokada):
                                            sala = check_sala(zad["Filia"], zad["Poziom"], d, st_m, en_m, grafik, sale_dane, gap_pref, lek["Lektor"])
                                            
                                            if sala:
                                                op_przyp, op_odp = "", ""
                                                moze_isc = True
                                                
                                                if zad["Typ"] == "Młodsza" and zad["Odbiór"]:
                                                    trasa = czas_tras.get(f"{zad['Odbiór']}_{zad['Filia']}", 15)
                                                    for op in opiekunki_dane:
                                                        if op["Filia"] == zad["Filia"] and any(b_s <= (st_m - trasa) and b_e >= st_m for (b_s, b_e) in op["Avail"][d]) and not any(g["Opiekunka"] == op["Imię"] and g["Dzień"] == d and is_overlap(st_m-trasa, st_m, g["Start"], g["End"], 0) for g in grafik_op):
                                                            op_przyp = op["Imię"]; break
                                                    for op in opiekunki_dane:
                                                        if op["Filia"] == zad["Filia"] and any(b_s <= en_m and b_e >= (en_m + trasa) for (b_s, b_e) in op["Avail"][d]) and not any(g["Opiekunka"] == op["Imię"] and g["Dzień"] == d and is_overlap(en_m, en_m+trasa, g["Start"], g["End"], 0) for g in grafik_op):
                                                            op_odp = op["Imię"]; break
                                                    
                                                    if not op_przyp or not op_odp: moze_isc = False
                                                
                                                if moze_isc:
                                                    op_str_format = ""
                                                    if op_przyp: 
                                                        grafik_op.extend([
                                                            {"Opiekunka": op_przyp, "Dzień": d, "Start": st_m-trasa, "End": st_m}, 
                                                            {"Opiekunka": op_odp, "Dzień": d, "Start": en_m, "End": en_m+trasa}
                                                        ])
                                                        c_p = f"{mins_to_time(st_m-trasa)}-{mins_to_time(st_m)}"
                                                        c_o = f"{mins_to_time(en_m)}-{mins_to_time(en_m+trasa)}"
                                                        o_odb = zad["Odbiór"]
                                                        op_str_format = f"\n🚶 Przyprowadza: {op_przyp} ({c_p}) z: {o_odb}\n🚶 Odprowadza: {op_odp} ({c_o}) do: {o_odb}"
                                                    
                                                    zaplanowane_dni.append({
                                                        "Grupa": zad["Grupa"], "Poziom": zad["Poziom"], "Filia": zad["Filia"],
                                                        "Dzień": d, "Start": st_m, "End": en_m, "Sala": sala, "Lektor": lek["Lektor"],
                                                        "Op_Str": op_str_format
                                                    })
                                                    dzien_ok = True
                                                    break 
                                    if not dzien_ok: break 
                                    
                                if len(zaplanowane_dni) == len(combo):
                                    grafik.extend(zaplanowane_dni)
                                    tracker[lek["Lektor"]]["grupy"].add(zad["Grupa"])
                                    tracker[lek["Lektor"]]["poziomy"].append(zad["Poziom"])
                                    
                                    if zad["Filia"] not in filia_day_load: filia_day_load[zad["Filia"]] = {}
                                    for d in combo:
                                        filia_day_load[zad["Filia"]][d] = filia_day_load[zad["Filia"]].get(d, 0) + 1
                                        if d not in tracker[lek["Lektor"]]["filie_dni"]: tracker[lek["Lektor"]]["filie_dni"][d] = set()
                                        tracker[lek["Lektor"]]["filie_dni"][d].add(zad["Filia"])
                                    return True
                    return False

                # FAZA 1: Rygorystyczna (bez migracji i z przerwami)
                znaleziono = proba_przypisania(strict_mode=True)
                
                # FAZA 2: Ratunkowa (zezwalaj na migrację, ignoruj limity VIP, zgódź się na 0 min przerwy)
                if not znaleziono:
                    znaleziono = proba_przypisania(strict_mode=False)

                if not znaleziono:
                    nieprzypisane.append({"Grupa": zad["Grupa"], "Filia": zad["Filia"], "Poziom": zad["Poziom"], "Problem": "Całkowity brak dyspozycyjności kadry/sal"})

            # === WYSZUKIWANIE OSTRZEŻEŃ ===
            warnings_zero = []
            warnings_min_grup = []
            warnings_max_grup = []
            warnings_rozstrzal = []
            
            for lek in lektorzy_dane:
                przypisane = len(tracker[lek["Lektor"]]["grupy"])
                
                if przypisane == 0:
                    warnings_zero.append(f"Lektor **{lek['Lektor']}** nie otrzymał **żadnej** grupy (0).")
                elif przypisane < lek["Limit_Min"]:
                    warnings_min_grup.append(f"Lektor **{lek['Lektor']}** chciał uczyć minimum {lek['Limit_Min']} grup, a ma tylko **{przypisane}**.")
                elif przypisane > lek["Limit_Max"]:
                    warnings_max_grup.append(f"Lektor **{lek['Lektor']}** prosił o maks {lek['Limit_Max']} grup, ale przydzielono mu aż **{przypisane}** (wymuszone brakiem kadr).")
                
                unikalne_poziomy = set(tracker[lek["Lektor"]]["poziomy"])
                if len(unikalne_poziomy) >= 4:
                    warnings_rozstrzal.append(f"Lektor **{lek['Lektor']}** ma {przypisane} grup, ale aż **{len(unikalne_poziomy)} różne poziomy** ({', '.join(sorted(unikalne_poziomy))}).")
            
            st.session_state.grafik = grafik
            st.session_state.nieprzypisane = nieprzypisane
            st.session_state.warnings_zero = warnings_zero
            st.session_state.warnings_min = warnings_min_grup
            st.session_state.warnings_max = warnings_max_grup
            st.session_state.warnings_lvl = warnings_rozstrzal
            st.session_state.wygenerowano = True

    # ==========================================
    # WIZUALIZACJA (ŁADOWANA Z PAMIĘCI SESSION_STATE)
    # ==========================================
    if st.session_state.get("wygenerowano", False):
        grafik = st.session_state.grafik
        nieprzypisane = st.session_state.nieprzypisane
        
        if nieprzypisane:
            st.error(f"🔴 Konflikty grafiku! ({len(nieprzypisane)} grup wylądowało w poczekalni).")
            st.dataframe(pd.DataFrame(nieprzypisane), use_container_width=True)
        else:
            st.success("🎉 Sukces! Przypisano wszystkie grupy!")
            
        if st.session_state.get("warnings_zero") or st.session_state.get("warnings_min") or st.session_state.get("warnings_lvl") or st.session_state.get("warnings_max"):
            with st.expander("⚠️ Raport ostrzeżeń: Niespełnione preferencje lektorów", expanded=True):
                if st.session_state.warnings_zero:
                    st.markdown("#### 🔴 Całkowity brak przydziału (0 grup)")
                    for w in st.session_state.warnings_zero:
                        st.write(f"- {w}")
                if st.session_state.warnings_min:
                    st.markdown("#### 📉 Brak wymaganej liczby grup (Poniżej MIN)")
                    for w in st.session_state.warnings_min:
                        st.write(f"- {w}")
                if st.session_state.warnings_max:
                    st.markdown("#### 📈 Przekroczenie etatu (Powyżej MAX)")
                    for w in st.session_state.warnings_max:
                        st.write(f"- {w}")
                if st.session_state.warnings_lvl:
                    st.markdown("#### 🔀 Duży rozstrzał poziomów")
                    for w in st.session_state.warnings_lvl:
                        st.write(f"- {w}")
            
        st.subheader("Wizualizacja Grafiku")
        
        widok_opcja = st.radio("Perspektywa:", ["Według Filii (Siatka 30-min)", "Według Lektorów (Indywidualnie)"], horizontal=True)
        df_grafik = pd.DataFrame(grafik)
        
        grafiki_do_eksportu = {}
        
        if widok_opcja == "Według Filii (Siatka 30-min)":
            tabs = st.tabs(filie_unikalne)
            for idx, tab in enumerate(tabs):
                f_nazwa = filie_unikalne[idx]
                with tab:
                    if not df_grafik.empty:
                        df_f = df_grafik[df_grafik["Filia"] == f_nazwa].copy()
                        df_grid = generate_grid_for_filia(df_f, f_nazwa, sale_dane)
                        
                        if not df_grid.empty:
                            uniq_leks = df_f["Lektor"].unique()
                            cmap = {lek: colors_pool[i % len(colors_pool)] for i, lek in enumerate(uniq_leks)}
                            
                            try: styled = df_grid.style.map(lambda x, c=cmap: style_cell_with_cmap(x, c)).format(format_cell_text)
                            except: styled = df_grid.style.applymap(lambda x, c=cmap: style_cell_with_cmap(x, c)).format(format_cell_text)
                            
                            grafiki_do_eksportu[f_nazwa] = styled
                            st.dataframe(styled, use_container_width=False, height=600)
                        else: st.info(f"Brak zajęć dla {f_nazwa}.")
                    else: st.info("Grafik pusty.")
            
            dl_filename = "Gotowe_Grafiki_Filie.xlsx"
                    
        else:
            if not df_grafik.empty:
                lektorzy_z_grafiku = sorted(df_grafik["Lektor"].unique())
                wybrany_lek = st.selectbox("Wybierz lektora do podglądu:", lektorzy_z_grafiku)
                
                df_lek = df_grafik[df_grafik["Lektor"] == wybrany_lek].copy()
                df_grid_lek = generate_grid_for_lektor(df_lek)
                
                if not df_grid_lek.empty:
                    uniq_g = df_lek["Grupa"].unique()
                    cmap_l = {g: colors_pool[i % len(colors_pool)] for i, g in enumerate(uniq_g)}
                    
                    try: styled_lek = df_grid_lek.style.map(lambda x, c=cmap_l: style_cell_with_cmap_lektor(x, c)).format(format_cell_text)
                    except: styled_lek = df_grid_lek.style.applymap(lambda x, c=cmap_l: style_cell_with_cmap_lektor(x, c)).format(format_cell_text)
                    
                    grafiki_do_eksportu[wybrany_lek] = styled_lek
                    st.dataframe(styled_lek, use_container_width=True, height=600)
                else: st.info("Grafik pusty.")
            else: st.info("Grafik pusty.")
            
            safe_lek_name = wybrany_lek.replace(" ", "_").replace("(", "").replace(")", "") if not df_grafik.empty else "Lektor"
            dl_filename = f"Grafik_{safe_lek_name}.xlsx"

        if grafiki_do_eksportu:
            st.markdown("---")
            excel_data = create_excel_download(grafiki_do_eksportu)
            st.download_button(
                label=f"📥 Pobierz wygenerowany grafik ({dl_filename})", 
                data=excel_data, 
                file_name=dl_filename,
                type="primary"
            )
